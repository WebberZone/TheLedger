"""Build the blank Accounts Tracker template bundled with the app.

Usage: build_blank_template.py <output.xlsx>

The original template is built from the owner's own spreadsheet and is
kept out of the repository (storage/app/templates/ is git-ignored). This
script builds a stand-in with the same structure populate_tracker.py
writes into: the same sheets, tables, header rows and Summary formulas,
with plain styling and no data. The export uses the private template when
one is installed and falls back to this one, so a fresh checkout (and CI)
can still produce a working tracker.

Layout contract with populate_tracker.py:

- Income/Expense/CF: one table each, columns B..M, header on row 3.
- Summary: Summary_Table, header on row 3, two sample rows. The first
  holds the cumulative ("everything up to this year") formulas and the
  second the per-year ones; the script copies them down.
- Balances: header on row 6, data from row 7, columns B..H, with the
  four snapshot columns D..G matching BALANCE_DATE_COLUMNS in
  TrackerExportController.
- PayPal: PayPal_USD with its header on row 6, PayPal_EUR below it on
  row 27. populate_tracker.py finds both by their table refs and moves
  the EUR section down when the USD rows need the space.
"""
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

OUTPUT = sys.argv[1]

DATE_FORMAT = 'dd mmmm yyyy'
MONEY_FORMAT = '#,##0.00'
GBP_FORMAT = '\\£#,##0.00'
FX_FORMAT = '0.0000'

TITLE_FONT = Font(bold=True, size=14)
HEADER_FONT = Font(bold=True, color='FFFFFF')
HEADER_FILL = PatternFill('solid', start_color='305496')
THIN = Side(style='thin', color='BFBFBF')
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

LEDGER_HEADERS = [
    ('Date', DATE_FORMAT, 16),
    ('Category', None, 16),
    ('Sub-Category', None, 24),
    ('Customer/Vendor', None, 24),
    ('Currency', None, 10),
    ('FX', FX_FORMAT, 10),
    ('Amount', MONEY_FORMAT, 12),
    ('VAT', MONEY_FORMAT, 12),
    ('Total', MONEY_FORMAT, 12),
    ('Amount GBP', GBP_FORMAT, 14),
    ('Description', None, 40),
    ('Payment Method', None, 18),
]

wb = Workbook()
wb.remove(wb.active)


def header_row(ws, row, first_col, headers):
    """Write a styled header row and format the row beneath it, which
    populate_tracker.py copies onto every row it writes."""
    for offset, (name, number_format, width) in enumerate(headers):
        col = first_col + offset
        cell = ws.cell(row=row, column=col, value=name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = CELL_BORDER

        sample = ws.cell(row=row + 1, column=col)
        sample.border = CELL_BORDER
        if number_format:
            sample.number_format = number_format

        letter = ws.cell(row=row, column=col).column_letter
        ws.column_dimensions[letter].width = width


def add_table(ws, name, ref):
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
    ws.add_table(table)


def title(ws, cell, text):
    ws[cell] = text
    ws[cell].font = TITLE_FONT


# Sheet order is part of the contract (TrackerExportTest checks it).
summary = wb.create_sheet('Summary')
ledgers = {name: wb.create_sheet(name) for name in ('Income', 'Expense', 'CF')}
balances = wb.create_sheet('Balances')
paypal = wb.create_sheet('PayPal')

# Ledger sheets -----------------------------------------------------------

for sheet_name, table_name, heading in (
    ('Income', 'Income_Table', 'Income'),
    ('Expense', 'Expenses_Table', 'Expenses'),
    ('CF', 'CF_Table', 'Cash flow'),
):
    ws = ledgers[sheet_name]
    title(ws, 'B1', heading)
    header_row(ws, 3, 2, LEDGER_HEADERS)
    add_table(ws, table_name, 'B3:M4')
    ws.freeze_panes = 'B4'

# Summary -----------------------------------------------------------------

title(summary, 'B1', 'Summary')
header_row(summary, 3, 2, [
    ('Year', '0', 10),
    ('Income', GBP_FORMAT, 16),
    ('Expenses', GBP_FORMAT, 16),
    ('Cash flow', GBP_FORMAT, 16),
    ('Profit', GBP_FORMAT, 16),
])
# Both sample rows need the same styling; header_row only formats one.
for col in range(2, 7):
    src = summary.cell(row=4, column=col)
    dst = summary.cell(row=5, column=col)
    dst.border = CELL_BORDER
    dst.number_format = src.number_format

YEAR = 'Summary_Table[[#This Row],[Year]]'
UP_TO_YEAR_END = f'"<="&DATE({YEAR},12,31)'
FROM_YEAR_START = f'">="&DATE({YEAR},1,1)'


def total(table, criteria):
    return f'SUMIFS({table}[Amount GBP],{table}[Date],' + f',{table}[Date],'.join(criteria) + ')'


PROFIT = '=Summary_Table[[#This Row],[Income]]-Summary_Table[[#This Row],[Expenses]]'

# Row 4: the earliest year rolls up all earlier history.
summary['C4'] = '=' + total('Income_Table', [UP_TO_YEAR_END])
summary['D4'] = '=' + total('Expenses_Table', [UP_TO_YEAR_END])
summary['E4'] = '=' + total('CF_Table', [UP_TO_YEAR_END])
summary['F4'] = PROFIT

# Row 5: every later year covers that calendar year only.
summary['C5'] = '=' + total('Income_Table', [FROM_YEAR_START, UP_TO_YEAR_END])
summary['D5'] = '=' + total('Expenses_Table', [FROM_YEAR_START, UP_TO_YEAR_END])
summary['E5'] = '=' + total('CF_Table', [FROM_YEAR_START, UP_TO_YEAR_END])
summary['F5'] = PROFIT

add_table(summary, 'Summary_Table', 'B3:F5')

# Balances ----------------------------------------------------------------

title(balances, 'B2', 'Account balances')
balances['B3'] = 'Balances at each date, in the account currency.'
header_row(balances, 6, 2, [
    ('Account', None, 22),
    ('Currency', None, 10),
    ('31 Dec 2023', MONEY_FORMAT, 14),
    ('31 Dec 2024', MONEY_FORMAT, 14),
    ('31 Dec 2025', MONEY_FORMAT, 14),
    ('17 Mar 2026', MONEY_FORMAT, 14),
    ('Notes', None, 40),
])

# PayPal ------------------------------------------------------------------

PAYPAL_HEADERS = [
    ('Date', DATE_FORMAT, 16),
    ('Description', None, 40),
    ('Source', None, 24),
    ('Amount', MONEY_FORMAT, 12),
    ('Balance', MONEY_FORMAT, 12),
]

title(paypal, 'B1', 'PayPal')
for table_name, heading, row in (('PayPal_USD', 'USD', 6), ('PayPal_EUR', 'EUR', 27)):
    paypal.cell(row=row - 1, column=2, value=heading).font = Font(bold=True)
    header_row(paypal, row, 2, PAYPAL_HEADERS)
    add_table(paypal, table_name, f'B{row}:F{row + 1}')

wb.save(OUTPUT)
print(f'Saved {OUTPUT}')
