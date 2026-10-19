"""Import transactions and balances from the legacy Accounts Tracker workbook.

Usage: import_tracker.py <path/to/Accounts Tracker.xlsx> [--db <path/to/database.sqlite>]

Reads the Income, Expense, CF and Balances sheets and inserts them into the
Ledger SQLite database. Categories, payment methods and the default business
unit are matched by name and created on demand. Rows are matched by the
natural key (date, counterparty, GBP amount) and skipped if already present.
"""
import argparse
import sqlite3
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

DEFAULT_DB_PATH = Path(__file__).resolve().parents[2] / "database" / "database.sqlite"

# sqlite3's deprecated default adapters are replaced by ISO strings.
sqlite3.register_adapter(date, lambda d: d.isoformat())
sqlite3.register_adapter(datetime, lambda d: d.isoformat())

# Normalise legacy spreadsheet aliases to the seeded payment-method names.
PAYMENT_METHOD_ALIASES = {
    "Paypal": "PayPal",
    "Amex": "Amex BA",
}


def to_date(value):
    """Excel/openpyxl date, datetime or ISO string -> date."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value)
        except ValueError:
            return None
    return None


def to_minor(value):
    """Float major amount -> integer minor units (pence/cents)."""
    if value is None:
        return 0
    if isinstance(value, str):
        value = value.strip()
        if value == "":
            return 0
        value = float(value)
    return int(round(value * 100))


def get_or_create_category(cur, name, type_, report_group):
    if name is None:
        return None
    name = str(name).strip()
    if name == "":
        return None

    cur.execute(
        "SELECT id FROM categories WHERE name = ? AND type = ?",
        (name, type_),
    )
    row = cur.fetchone()
    if row:
        return row[0]

    cur.execute(
        "INSERT INTO categories (name, type, report_group, is_active, sort_order, created_at, updated_at) "
        "VALUES (?, ?, ?, 1, 0, ?, ?)",
        (name, type_, report_group, datetime.now(), datetime.now()),
    )
    return cur.lastrowid


def get_or_create_payment_method(cur, name):
    if name is None:
        return None
    name = str(name).strip()
    if name == "":
        return None

    name = PAYMENT_METHOD_ALIASES.get(name, name)

    cur.execute("SELECT id FROM payment_methods WHERE name = ?", (name,))
    row = cur.fetchone()
    if row:
        return row[0]

    cur.execute(
        "INSERT INTO payment_methods (name, is_personal, is_active, created_at, updated_at) "
        "VALUES (?, 0, 1, ?, ?)",
        (name, datetime.now(), datetime.now()),
    )
    return cur.lastrowid


def get_or_create_business_unit(cur, name):
    cur.execute("SELECT id FROM business_units WHERE name = ?", (name,))
    row = cur.fetchone()
    if row:
        return row[0]

    cur.execute(
        "INSERT INTO business_units (name, is_active, created_at, updated_at) "
        "VALUES (?, 1, ?, ?)",
        (name, datetime.now(), datetime.now()),
    )
    return cur.lastrowid


def transaction_exists(cur, type_, occurred_on, counterparty, amount_gbp_minor):
    """Check whether a transaction with the same natural key is already imported."""
    if counterparty is None:
        cur.execute(
            "SELECT id FROM transactions WHERE type = ? AND occurred_on = ? "
            "AND counterparty IS NULL AND amount_gbp_minor = ? LIMIT 1",
            (type_, occurred_on, amount_gbp_minor),
        )
    else:
        cur.execute(
            "SELECT id FROM transactions WHERE type = ? AND occurred_on = ? "
            "AND counterparty = ? AND amount_gbp_minor = ? LIMIT 1",
            (type_, occurred_on, counterparty, amount_gbp_minor),
        )
    return cur.fetchone() is not None


def import_ledger(cur, ws, type_, business_unit_id, counts):
    """Import rows from a ledger sheet (Income, Expense, CF)."""
    # Header row is row 21; data starts at row 22.
    for row in ws.iter_rows(min_row=22, max_row=ws.max_row):
        date_cell = row[1]  # B
        category_cell = row[2]  # C
        subcategory_cell = row[3]  # D
        counterparty_cell = row[4]  # E
        currency_cell = row[5]  # F
        fx_cell = row[6]  # G
        amount_cell = row[7]  # H
        vat_cell = row[8]  # I
        gbp_cell = row[10]  # K
        description_cell = row[11]  # L
        payment_method_cell = row[12]  # M

        occurred_on = to_date(date_cell.value)
        if occurred_on is None:
            continue

        category = str(category_cell.value).strip() if category_cell.value else None
        subcategory = str(subcategory_cell.value).strip() if subcategory_cell.value else None
        counterparty = str(counterparty_cell.value).strip() if counterparty_cell.value else None
        currency = str(currency_cell.value).strip().upper() if currency_cell.value else "GBP"
        fx_rate = float(fx_cell.value) if fx_cell.value is not None else 1.0
        amount_minor = to_minor(amount_cell.value)
        vat_minor = to_minor(vat_cell.value)
        amount_gbp_minor = to_minor(gbp_cell.value)
        description = str(description_cell.value).strip() if description_cell.value else None
        payment_method = str(payment_method_cell.value).strip() if payment_method_cell.value else None

        category_id = get_or_create_category(cur, subcategory, type_, category)
        payment_method_id = get_or_create_payment_method(cur, payment_method)

        if transaction_exists(cur, type_, occurred_on, counterparty, amount_gbp_minor):
            counts.setdefault("skipped", 0)
            counts["skipped"] += 1
            continue

        cur.execute(
            "INSERT INTO transactions (type, occurred_on, counterparty, description, "
            "amount_minor, vat_minor, currency, fx_rate, amount_gbp_minor, category_id, "
            "business_unit_id, payment_method_id, is_billable, billed_ref, is_draft, notes, "
            "ocr_json, created_at, updated_at) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0, NULL, 0, NULL, NULL, ?, ?)",
            (
                type_,
                occurred_on,
                counterparty,
                description,
                amount_minor,
                vat_minor,
                currency,
                fx_rate,
                amount_gbp_minor,
                category_id,
                business_unit_id,
                payment_method_id,
                datetime.now(),
                datetime.now(),
            ),
        )
        counts[type_] += 1


def parse_header_date(value):
    """Parse strings like '31 Dec 2023' or '17 Mar 2026 (current)'."""
    if value is None:
        return None
    text = str(value).strip()
    if text == "":
        return None

    # Remove parenthetical notes.
    text = text.split("(")[0].strip()

    try:
        return datetime.strptime(text, "%d %b %Y").date()
    except ValueError:
        return None


def import_balances(cur, ws, counts):
    """Import rows from the Balances sheet."""
    # Header row is row 6; data starts at row 7.
    header = list(ws.iter_rows(min_row=6, max_row=6))[0]
    date_columns = []
    # Header: B=Account, C=Currency, D-G=dates. Skip A/B/C.
    for idx, cell in enumerate(header, start=1):
        if idx <= 3:
            continue
        parsed = parse_header_date(cell.value)
        if parsed:
            date_columns.append((idx, parsed))

    for row in ws.iter_rows(min_row=7, max_row=ws.max_row):
        if len(row) < 3:
            continue

        account = row[1].value  # B
        currency = row[2].value  # C
        notes = row[7].value if len(row) > 7 else None  # H

        if account is None or currency is None:
            continue

        account = str(account).strip()
        currency = str(currency).strip().upper()
        notes = str(notes).strip() if notes else None

        for col_idx, balance_on in date_columns:
            if col_idx > len(row):
                continue
            value = row[col_idx - 1].value
            if value is None or value == "":
                continue

            balance_minor = to_minor(value)

            cur.execute(
                "SELECT id FROM account_balances WHERE account = ? AND currency = ? AND balance_on = ? LIMIT 1",
                (account, currency, balance_on),
            )
            if cur.fetchone():
                counts["skipped"] += 1
                continue

            cur.execute(
                "INSERT INTO account_balances (account, currency, balance_on, balance_minor, notes, created_at, updated_at) "
                "VALUES (?, ?, ?, ?, ?, ?, ?)",
                (account, currency, balance_on, balance_minor, notes, datetime.now(), datetime.now()),
            )
            counts["balances"] += 1


def main():
    parser = argparse.ArgumentParser(
        description="Import the legacy Accounts Tracker workbook into Ledger."
    )
    parser.add_argument("workbook", help="Path to the Accounts Tracker.xlsx file")
    parser.add_argument(
        "--db",
        default=DEFAULT_DB_PATH,
        help="Path to the SQLite database (default: %(default)s)",
    )
    args = parser.parse_args()

    workbook_path = Path(args.workbook)
    if not workbook_path.exists():
        print(f"File not found: {workbook_path}", file=sys.stderr)
        sys.exit(1)

    db_path = Path(args.db)

    wb = openpyxl.load_workbook(workbook_path, data_only=True)

    conn = sqlite3.connect(db_path)
    conn.execute("PRAGMA foreign_keys = ON")
    cur = conn.cursor()

    counts = {"income": 0, "expense": 0, "cash_flow": 0, "balances": 0, "skipped": 0}

    business_unit_id = get_or_create_business_unit(cur, "Beyond Develop")

    try:
        import_ledger(cur, wb["Income"], "income", business_unit_id, counts)
        import_ledger(cur, wb["Expense"], "expense", business_unit_id, counts)
        import_ledger(cur, wb["CF"], "cash_flow", business_unit_id, counts)

        if "Balances" in wb.sheetnames:
            import_balances(cur, wb["Balances"], counts)

        conn.commit()
        print(
            f"Imported {counts['income']} income, {counts['expense']} expense, "
            f"{counts['cash_flow']} cash-flow transactions and {counts['balances']} balance snapshots."
            f" Skipped {counts['skipped']} duplicates."
        )
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


if __name__ == "__main__":
    main()
